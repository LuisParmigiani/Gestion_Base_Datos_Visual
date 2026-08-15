"""Excel del banco → tabla `movimiento_banco`.

    ia/.venv/bin/python -m portal.importar ../TRANSFERENCIAS.xlsx

Porteo de `web/src/lib/importar.ts`, con una simplificación a propósito:

  - **No hay tabla `cuenta`.** La cuenta se guarda como texto: la elige el dueño
    al subir, o sale del nombre de cada hoja. En `web/` una hoja desconocida era
    un error; acá no hay contra qué cotejarla, y una abstracción con una sola
    implementación se inlinea.

Sí hay **tabla de cargas** (`carga_extracto`): una fila por archivo subido, para
poder mirar las cargas viejas y **deshacer** la que se subió por error.

Lo que NO se simplificó: la huella, que es lo que hace **seguro reimportar** el
mismo archivo, y la columna `SALDO`, que **no se lee** — es información del
banco, no del proyecto.
"""

import hashlib
import sys
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert

from portal.db import Sesion
from portal.extracto import (
    extraer_coelsa,
    extraer_cuit,
    extraer_tipo,
    fecha_de_celda,
    huella,
    parsear_importe,
)
from portal.modelos import CargaExtracto, MovimientoBanco


def _texto(valor) -> str:
    """La celda como texto, sin inventar decimales en los números enteros."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def importar(
    origen: Path | bytes,
    cuenta: str | None = None,
    nombre_archivo: str | None = None,
    sesion=None,
) -> CargaExtracto:
    """Importa un Excel y devuelve la carga, con su resumen por hoja adentro.

    `cuenta` fuerza a qué cuenta van TODAS las filas. Sirve para los archivos que
    traen una sola cuenta con un nombre de hoja cualquiera (`Hoja1`). Si no se
    pasa, cada hoja va a la cuenta que lleva su nombre — que es como viene el
    archivo con varias cuentas adentro.

    ⚠️ La cuenta entra en la huella. Subir el mismo archivo a dos cuentas
    distintas mete las filas dos veces, y está bien: son cuentas distintas.
    """
    bytes_ = origen.read_bytes() if isinstance(origen, Path) else origen
    libro = load_workbook(BytesIO(bytes_), data_only=True, read_only=True)

    propia = sesion is None
    sesion = sesion or Sesion()
    try:
        carga = CargaExtracto(
            archivo=nombre_archivo or (origen.name if isinstance(origen, Path) else "subido.xlsx"),
            hash_archivo=hashlib.sha256(bytes_).hexdigest(),
            cuenta_forzada=cuenta,
        )
        sesion.add(carga)
        sesion.flush()  # necesita id antes de que los movimientos lo referencien

        resumen = [_importar_hoja(sesion, hoja, carga.id, cuenta) for hoja in libro.worksheets]
        carga.resumen = {"hojas": resumen}
        carga.filas_leidas = sum(r["leidas"] for r in resumen)
        carga.filas_nuevas = sum(r["nuevas"] for r in resumen)
        carga.filas_duplicadas = sum(r["duplicadas"] for r in resumen)

        if propia:
            sesion.commit()
            sesion.refresh(carga)
        return carga
    finally:
        libro.close()
        if propia:
            sesion.close()


def ya_subido(sesion, contenido: bytes) -> CargaExtracto | None:
    """La carga anterior de este MISMO archivo, si la hubo. Para poder avisar."""
    return sesion.scalars(
        select(CargaExtracto)
        .where(CargaExtracto.hash_archivo == hashlib.sha256(contenido).hexdigest())
        .order_by(CargaExtracto.id.desc())
    ).first()


def _importar_hoja(sesion, hoja, carga_id: int, cuenta_forzada: str | None) -> dict:
    # La cuenta que eligió el dueño gana sobre el nombre de la hoja.
    cuenta = cuenta_forzada or hoja.title
    filas, sin_fecha, ambiguas = [], 0, 0
    vistas = set()  # la primera aparición gana, igual que el ON CONFLICT

    for numero, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
        if numero == 1:
            continue  # encabezado
        # Columnas: FECHA · NRO TRANSACCION · DESCRIPCION · IMPORTE · SALDO.
        # SALDO no se lee, ni siquiera para descartarlo después.
        celda_fecha = fila[0] if len(fila) > 0 else None
        nro = _texto(fila[1] if len(fila) > 1 else None)
        descripcion = _texto(fila[2] if len(fila) > 2 else None)
        crudo = _texto(fila[3] if len(fila) > 3 else None)

        if not celda_fecha and not nro and not descripcion and not crudo:
            continue  # fila vacía

        fecha = fecha_de_celda(celda_fecha)
        if fecha is None:
            sin_fecha += 1
            continue  # sin fecha no hay movimiento

        importe = parsear_importe(crudo)
        ambiguas += importe.ambiguo
        marca = huella(
            cuenta=cuenta,
            fecha=fecha,
            nro=nro,
            importe_crudo=crudo,
            descripcion=descripcion,
        )
        if marca in vistas:
            continue
        vistas.add(marca)

        filas.append({
            "carga_id": carga_id,
            "cuenta": cuenta,
            "fecha": fecha,
            "nro_transaccion": nro or None,
            "descripcion": descripcion,
            "importe_crudo": crudo,
            "importe": Decimal(importe.importe) if importe.importe else None,
            "es_ingreso": importe.es_ingreso,
            "requiere_revision": importe.ambiguo,
            "cuit_emisor": extraer_cuit(descripcion),
            "coelsa_id": extraer_coelsa(descripcion),
            "tipo": extraer_tipo(descripcion),
            "huella": marca,
            "hoja": hoja.title,
            "fila": numero,
        })

    nuevas = 0
    if filas:
        # Un solo INSERT con ON CONFLICT DO NOTHING: la huella UNIQUE es la que
        # deduplica, así reimportar el mismo Excel entra 0 filas sin romper nada.
        #
        # Va contra la tabla (`__table__`) y no contra la clase, y cuenta con
        # RETURNING en vez de `rowcount`: con ON CONFLICT, `rowcount` no
        # distingue lo insertado de lo ignorado. Las filas que vuelven son
        # exactamente las que entraron.
        resultado = sesion.execute(
            insert(MovimientoBanco.__table__)
            .on_conflict_do_nothing(index_elements=["huella"])
            .returning(MovimientoBanco.__table__.c.id),
            filas,
        )
        nuevas = len(resultado.all())

    fechas = [f["fecha"] for f in filas]
    return {
        "cuenta": cuenta,
        "hoja": hoja.title,
        "leidas": len(filas) + sin_fecha,
        "nuevas": nuevas,
        "duplicadas": len(filas) - nuevas,
        "a_revisar": ambiguas,
        "sin_fecha": sin_fecha,
        "fecha_min": min(fechas) if fechas else None,
        "fecha_max": max(fechas) if fechas else None,
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("uso: python -m portal.importar <archivo.xlsx> [cuenta]")

    carga = importar(Path(sys.argv[1]), cuenta=sys.argv[2] if len(sys.argv) > 2 else None)
    for linea in carga.resumen["hojas"]:
        print(
            f"  {linea['cuenta']:<22} leídas {linea['leidas']:>5} · "
            f"nuevas {linea['nuevas']:>5} · duplicadas {linea['duplicadas']:>5} · "
            f"a revisar {linea['a_revisar']:>3} · sin fecha {linea['sin_fecha']:>3} · "
            f"{linea['fecha_min']} → {linea['fecha_max']}"
        )
