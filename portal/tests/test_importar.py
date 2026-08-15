"""Pruebas del importador del Excel. Los libros se FABRICAN acá, en memoria.

Ningún extracto real entra en los tests (regla del proyecto). Los casos cubren
lo que en `web/` costó sorpresas: los dos formatos de número mezclados, la
columna SALDO que no se lee, y que reimportar el mismo archivo no duplique.
"""

import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO

from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select

from portal.app import app

from portal.db import Sesion, motor
from portal.importar import importar as importar_archivo
from portal.modelos import (
    CONFIRMADO,
    PENDIENTE,
    Base,
    CargaExtracto,
    Cliente,
    Comprobante,
    MovimientoBanco,
    MovimientoCuenta,
    registrar_movimiento,
)


def importar(origen, **kw):
    """El resumen por hoja, que es lo que mira casi todo el archivo.

    `importar_archivo` devuelve la carga entera —hace falta para las pruebas de
    cargas— pero la mayoría de los casos solo mira las hojas.
    """
    return importar_archivo(origen, **kw).resumen["hojas"]

if not str(motor.url).endswith("portal_test"):
    raise SystemExit(
        f"Los tests apuntan a {motor.url!r}. Falta PORTAL_DATABASE_URL con portal_test."
    )

ENCABEZADO = ["FECHA", "NRO TRANSACCION", "DESCRIPCION", "IMPORTE", "SALDO"]


def libro(filas, hoja="SUCESION") -> bytes:
    """Un .xlsx en memoria con el encabezado y las filas que se le pasen."""
    wb = Workbook()
    wb.active.title = hoja
    wb.active.append(ENCABEZADO)
    for fila in filas:
        wb.active.append(fila)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class Importacion(unittest.TestCase):
    def setUp(self):
        Base.metadata.drop_all(motor)
        Base.metadata.create_all(motor)

    def movimientos(self):
        with Sesion() as sesion:
            return sesion.scalars(
                select(MovimientoBanco).order_by(MovimientoBanco.id)
            ).all()

    def test_importa_una_fila_con_todos_sus_campos_derivados(self):
        importar(libro([
            [date(2026, 2, 27), "123456", "ING TRANSF 27123456789 GOMEZ/A", "$ 368.500,00", "$ 999"],
        ]))
        movimiento = self.movimientos()[0]
        self.assertEqual(movimiento.importe, Decimal("368500.00"))
        self.assertTrue(movimiento.es_ingreso)
        self.assertFalse(movimiento.requiere_revision)
        self.assertEqual(movimiento.cuit_emisor, "27123456789")
        self.assertEqual(movimiento.tipo, "ING TRANSF")
        self.assertEqual(movimiento.cuenta, "SUCESION")
        self.assertEqual(movimiento.fila, 2)

    def test_la_columna_saldo_no_se_guarda_en_ningun_lado(self):
        # Es información del banco, no del proyecto. Ni siquiera se lee.
        importar(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", "$ 7654321,99"],
        ]))
        movimiento = self.movimientos()[0]
        self.assertNotIn("7654321", str(vars(movimiento)))

    def test_reimportar_el_mismo_archivo_no_duplica_nada(self):
        bytes_ = libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
            [date(2026, 2, 28), "2", "TEF 20111111112 Y", "-$ 50,00", ""],
        ])
        primero = importar(bytes_)
        segundo = importar(bytes_)

        self.assertEqual(primero[0]["nuevas"], 2)
        self.assertEqual(segundo[0]["nuevas"], 0)
        self.assertEqual(segundo[0]["duplicadas"], 2)
        with Sesion() as sesion:
            self.assertEqual(
                sesion.scalar(select(func.count()).select_from(MovimientoBanco)), 2
            )

    def test_un_importe_ambiguo_entra_sin_importe_y_marcado(self):
        # '$ 1.500' puede ser 1500 o 1,50. No se adivina: se marca.
        resumen = importar(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 1.500", ""],
        ]))
        movimiento = self.movimientos()[0]
        self.assertIsNone(movimiento.importe)
        self.assertTrue(movimiento.requiere_revision)
        self.assertEqual(movimiento.importe_crudo, "$ 1.500")
        self.assertEqual(resumen[0]["a_revisar"], 1)

    def test_una_fila_sin_fecha_no_entra_pero_se_cuenta(self):
        resumen = importar(libro([
            ["", "1", "ALGO SIN FECHA", "$ 100,00", ""],
            [date(2026, 2, 27), "2", "ING TRANSF 27123456789 X", "$ 200,00", ""],
        ]))
        self.assertEqual(resumen[0]["sin_fecha"], 1)
        self.assertEqual(resumen[0]["nuevas"], 1)
        self.assertEqual(resumen[0]["leidas"], 2)

    def test_las_filas_vacias_se_saltean_sin_contarse(self):
        resumen = importar(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
            [None, None, None, None, None],
        ]))
        self.assertEqual(resumen[0]["leidas"], 1)

    def test_cada_hoja_es_una_cuenta_distinta(self):
        wb = Workbook()
        for nombre in ["SUCESION", "OTRA CUENTA"]:
            hoja = wb.create_sheet(nombre)
            hoja.append(ENCABEZADO)
            hoja.append([date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""])
        del wb["Sheet"]
        buffer = BytesIO()
        wb.save(buffer)

        resumen = importar(buffer.getvalue())
        self.assertEqual([r["cuenta"] for r in resumen], ["SUCESION", "OTRA CUENTA"])
        # Misma fila, distinta cuenta: la huella incluye la cuenta, así que son dos.
        self.assertEqual(sum(r["nuevas"] for r in resumen), 2)

    def test_el_resumen_trae_el_rango_de_fechas(self):
        resumen = importar(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
            [date(2026, 1, 5), "2", "TEF 20111111112 Y", "$ 200,00", ""],
        ]))
        self.assertEqual(resumen[0]["fecha_min"], "2026-01-05")
        self.assertEqual(resumen[0]["fecha_max"], "2026-02-27")


if __name__ == "__main__":
    unittest.main()


class Cuentas(unittest.TestCase):
    """Un archivo por cuenta: el dueño elige a cuál van las filas."""

    def setUp(self):
        Base.metadata.drop_all(motor)
        Base.metadata.create_all(motor)

    def test_sin_elegir_cuenta_cada_hoja_va_a_la_suya(self):
        resumen = importar(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
        ], hoja="SUCESION"))
        self.assertEqual(resumen[0]["cuenta"], "SUCESION")

    def test_eligiendo_cuenta_todas_las_filas_van_ahi(self):
        # El caso real: un Excel de una sola cuenta cuya hoja se llama "Hoja1".
        resumen = importar(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
        ], hoja="Hoja1"), cuenta="BRUBANK")
        self.assertEqual(resumen[0]["cuenta"], "BRUBANK")
        self.assertEqual(resumen[0]["hoja"], "Hoja1")
        with Sesion() as sesion:
            self.assertEqual(sesion.scalars(select(MovimientoBanco)).one().cuenta, "BRUBANK")

    def test_el_mismo_archivo_en_dos_cuentas_entra_dos_veces(self):
        # No es un duplicado: son movimientos de cuentas distintas. La cuenta
        # forma parte de la huella justamente por esto.
        bytes_ = libro([[date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""]])
        importar(bytes_, cuenta="UNA")
        importar(bytes_, cuenta="OTRA")
        with Sesion() as sesion:
            self.assertEqual(
                sesion.scalar(select(func.count()).select_from(MovimientoBanco)), 2)


class Cargas(unittest.TestCase):
    """Cada subida queda registrada, para poder mirarla y para poder deshacerla."""

    def setUp(self):
        Base.metadata.drop_all(motor)
        Base.metadata.create_all(motor)

    def test_cada_subida_deja_su_carga_con_los_conteos(self):
        bytes_ = libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
            [date(2026, 2, 28), "2", "TEF 20111111112 Y", "-$ 50,00", ""],
        ])
        carga = importar_archivo(bytes_, nombre_archivo="febrero.xlsx")

        self.assertEqual(carga.archivo, "febrero.xlsx")
        self.assertEqual(carga.filas_nuevas, 2)
        self.assertEqual(carga.filas_duplicadas, 0)
        self.assertEqual(len(carga.hash_archivo), 64)

    def test_cada_movimiento_sabe_de_qué_carga_vino(self):
        carga = importar_archivo(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
        ]))
        with Sesion() as sesion:
            self.assertEqual(
                sesion.scalars(select(MovimientoBanco)).one().carga_id, carga.id)

    def test_subir_el_mismo_archivo_se_detecta_por_su_hash(self):
        from portal.importar import ya_subido
        bytes_ = libro([[date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""]])
        primera = importar_archivo(bytes_, nombre_archivo="mismo.xlsx")
        with Sesion() as sesion:
            encontrada = ya_subido(sesion, bytes_)
            self.assertEqual(encontrada.id, primera.id)
            self.assertIsNone(ya_subido(sesion, libro([])))


class DeshacerCarga(unittest.TestCase):
    """Deshacer una carga no puede dejar un saldo apoyado en una fila que ya no existe."""

    def setUp(self):
        Base.metadata.drop_all(motor)
        Base.metadata.create_all(motor)
        self.cliente_http = TestClient(app)

    def _carga_con_una_fila(self):
        return importar_archivo(libro([
            [date(2026, 2, 27), "1", "ING TRANSF 27123456789 X", "$ 100,00", ""],
        ]), nombre_archivo="equivocado.xlsx")

    def test_deshacer_una_carga_sin_acreditar_borra_sus_movimientos(self):
        carga = self._carga_con_una_fila()
        respuesta = self.cliente_http.post(
            f"/admin/extracto/carga/{carga.id}/deshacer", follow_redirects=False)

        self.assertEqual(respuesta.status_code, 303)
        with Sesion() as sesion:
            self.assertEqual(
                sesion.scalar(select(func.count()).select_from(MovimientoBanco)), 0)
            self.assertIsNone(sesion.get(CargaExtracto, carga.id))

    def test_deshacer_anula_lo_acreditado_y_devuelve_el_comprobante_al_lote(self):
        carga = self._carga_con_una_fila()
        with Sesion() as sesion:
            cliente = Cliente(nombre="Ana", apellido="Test", email="a@test.test")
            sesion.add(cliente)
            sesion.flush()
            comprobante = Comprobante(cliente_id=cliente.id, archivo="x.jpg",
                                      estado=CONFIRMADO, confianza=90)
            sesion.add(comprobante)
            sesion.flush()
            fila = sesion.scalars(select(MovimientoBanco)).one()
            fila.comprobante_id = comprobante.id
            registrar_movimiento(sesion, cliente, date(2026, 2, 27), "Transferencia",
                                 Decimal("100.00"), comprobante)
            sesion.commit()
            self.assertEqual(sesion.get(Cliente, cliente.id).saldo, Decimal("100.00"))

        # La pantalla de confirmación tiene que AVISAR antes de borrar.
        aviso = self.cliente_http.get(f"/admin/extracto/carga/{carga.id}/deshacer")
        self.assertIn("ya acreditaron plata", aviso.text)

        self.cliente_http.post(f"/admin/extracto/carga/{carga.id}/deshacer",
                               follow_redirects=False)

        with Sesion() as sesion:
            # El saldo vuelve a cero, y las DOS líneas quedan en el libro.
            self.assertEqual(sesion.scalars(select(Cliente)).one().saldo, Decimal("0.00"))
            self.assertEqual(
                sesion.scalar(select(func.count()).select_from(MovimientoCuenta)), 2)
            comprobante = sesion.scalars(select(Comprobante)).one()
            self.assertEqual(comprobante.estado, PENDIENTE)
            self.assertIsNone(comprobante.confianza)
            self.assertIn("Se deshizo la carga", comprobante.motivo)
